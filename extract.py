import openpyxl
import json

wb = openpyxl.load_workbook('등장인물.xlsx')
data = {}
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    sheet_data = []
    for row in ws.iter_rows(values_only=True):
        sheet_data.append([str(cell) if cell is not None else '' for cell in row])
    data[sheet_name] = sheet_data

with open('excel_data.json', 'w', encoding='utf-8') as f:
    json.dump(data, f, ensure_ascii=False, indent=2)
